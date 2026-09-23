#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generic PCIbex / PennController post-SPR event timing extractor
Version 0.2.0

Purpose
-------
For each reconstructed trial containing self-paced reading (SPR), locate the
FINAL SPR segment and then collect EVERY logged event that occurs after that
final segment within the same trial.

Unlike the original project-specific script, this version:
- has no hard-coded condition name;
- has no hard-coded subject-type variable;
- does not assume the post-SPR event is a Scale judgment;
- automatically retains non-personal custom variables;
- automatically identifies time-like custom fields;
- anonymises participants as Participant_Number;
- does not export MD5, reception time, names, emails, or common personal fields.

Output workbook
---------------
1. Trials
   One row per SPR trial, with:
   - anonymous participant number;
   - label/order/sentence;
   - automatically detected experimental variables;
   - final SPR segment information;
   - number of post-SPR events;
   - first and last post-SPR event times;
   - elapsed time from the final SPR EventTime to first/last post-SPR event;
   - an alternative elapsed time from estimated final segment END
     (Final_SPR_EventTime + Final_SPR_RT).

2. Post_SPR_Events
   Long format: one row for EVERY event logged after the final SPR segment.
   Includes event labels, values, EventTime, elapsed time from the final SPR
   event, elapsed time from estimated SPR end, and all detected custom variables.

3. Time_Like_Fields
   Long format: every numeric custom field whose NAME looks time-related
   (e.g. Reading Time, RT, latency, duration, response time, judgment time).
   These fields are discovered automatically.

4. Variables
   Lists retained and privacy-excluded custom variables.

5. Diagnostic
   Compares gaps between consecutive SPR EventTimes with logged segment RTs.
   This helps determine whether DashedSentence EventTime most likely marks the
   START or END of each segment.

Dependencies
------------
    pip install openpyxl

Usage
-----
    python post_spr_timing.py raw_results.csv
    python post_spr_timing.py raw_results.csv output.xlsx
    python post_spr_timing.py --version

Important interpretation note
-----------------------------
PCIbex/PennController EventTime conventions can vary by controller/version.
Therefore this program reports BOTH:
A) delta from Final_SPR_EventTime
B) delta from estimated final SPR end = Final_SPR_EventTime + Final_SPR_RT

The Diagnostic sheet helps decide which interpretation is better supported by
the actual raw data.
"""

from __future__ import annotations

import argparse
import csv
import re
import sys
import unicodedata
from collections import OrderedDict, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "0.2.0"
DEFAULT_OUTPUT_NAME = "post_spr_timing.xlsx"

BASE_FIELDS = [
    "ReceptionTime", "MD5", "Controller", "OrderNumber", "InnerNumber",
    "Label", "LatinSquareGroup", "PennElementType", "PennElementName",
    "Parameter", "Value", "EventTime",
]

COMMENT_PATTERN = re.compile(r"^#\s*(\d+)\.\s*(.+?)\.?\s*$")
EMPTY_TOKENS = {"", "NULL", "null", "NA", "N/A", None}

SPR_KEYWORDS = ("dashedsentence", "dashed sentence")

PII_EXACT = {
    "name", "nome", "full name", "fullname",
    "email", "e mail", "e-mail", "mail",
    "phone", "phone number", "telephone", "telefone", "telemovel",
    "mobile", "mobile phone", "whatsapp",
    "participant id", "participant_id", "participantid",
    "subject id", "subject_id", "subjectid",
    "student id", "student_id", "studentid",
    "worker id", "worker_id", "workerid",
    "prolific id", "prolific_id", "prolificid", "prolific pid",
    "sona id", "sona_id", "sonaid",
    "matricula", "numero de estudante", "numero estudante",
    "age", "idade", "gender", "genero", "género", "sex", "sexo",
    "school", "escola", "institution", "instituicao", "instituição",
    "course", "curso", "nationality", "nacionalidade",
    "country", "pais", "país",
}

PII_CONTAINS = (
    "email address",
    "endereco de email",
    "endereço de email",
    "phone number",
    "numero de telefone",
    "número de telefone",
    "ip address",
    "ip_address",
)

TIME_FIELD_KEYWORDS = (
    "time", "tempo", "rt", "latency", "latencia", "latência",
    "duration", "duracao", "duração", "response time", "reaction time",
    "reading time", "judgment time", "judgement time", "click time",
    "elapsed", "timestamp",
)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def strip_accents(value: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFKD", value)
        if not unicodedata.combining(c)
    )


def norm(value) -> str:
    if value is None:
        return ""
    return strip_accents(str(value)).lower().strip()


def normalised_field_name(name: str) -> str:
    n = norm(name).replace("_", " ").replace("-", " ")
    return re.sub(r"\s+", " ", n).strip()


def clean_col_name(name: str) -> str:
    name = re.sub(r"\(.*?\)", "", name)
    return name.replace("?", "").strip()


def decode_text(value):
    if value is None:
        return value
    return (
        str(value)
        .replace("%2C", ",")
        .replace("%27", "'")
        .replace("%22", '"')
    )


def is_empty(value) -> bool:
    return value in EMPTY_TOKENS or norm(value) in {"", "null", "na", "n/a"}


def is_pii_field(name: str) -> bool:
    n = normalised_field_name(name)
    exact = {normalised_field_name(x) for x in PII_EXACT}
    if n in exact:
        return True
    return any(normalised_field_name(x) in n for x in PII_CONTAINS)


def is_sentence_like(name: str) -> bool:
    n = normalised_field_name(name)
    return (
        n == "sentence"
        or n == "frase"
        or n.startswith("sentence ")
        or n.startswith("frase ")
        or "sentence text" in n
        or "texto da frase" in n
    )


def is_reading_time_like(name: str) -> bool:
    n = normalised_field_name(name)
    return (
        "reading time" in n
        or "tempo de leitura" in n
        or n in {"rt", "rt ms", "readingtime"}
    )


def is_time_like_field(name: str) -> bool:
    n = normalised_field_name(name)
    # Treat isolated "rt" as a token rather than matching arbitrary words.
    tokens = set(n.split())
    if "rt" in tokens:
        return True
    return any(normalised_field_name(k) in n for k in TIME_FIELD_KEYWORDS if k != "rt")


def safe_int(value):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        try:
            return int(float(str(value).strip()))
        except (TypeError, ValueError):
            return None


def safe_number(value):
    text = str(value).strip()
    try:
        if re.fullmatch(r"[-+]?\d+", text):
            return int(text)
        return float(text)
    except (TypeError, ValueError):
        return None


def event_is_spr(record) -> bool:
    combined = " ".join([
        norm(record.get("Controller")),
        norm(record.get("PennElementType")),
        norm(record.get("PennElementName")),
    ])
    return any(keyword in combined for keyword in SPR_KEYWORDS)


def session_key(record):
    # Internal use only; raw values are never exported.
    md5 = record.get("MD5", "")
    reception = record.get("ReceptionTime", "")
    if md5 or reception:
        return (md5, reception)

    return (
        record.get("Controller", ""),
        record.get("EventTime", ""),
        record.get("OrderNumber", ""),
    )


# ---------------------------------------------------------------------------
# Parse raw PCIbex/PennController file
# ---------------------------------------------------------------------------

def parse_results_file(path: Path):
    column_map = {}
    pending_block = OrderedDict()
    rows = []
    raw_index = 0

    def flush_block():
        nonlocal pending_block
        for number, name in pending_block.items():
            column_map[number] = clean_col_name(name)
        pending_block = OrderedDict()

    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
        for raw_line in f:
            line = raw_line.rstrip("\r\n")
            if not line.strip():
                continue

            if line.lstrip().startswith("#"):
                match = COMMENT_PATTERN.match(line.strip())
                if match:
                    number = int(match.group(1))
                    if number >= 13:
                        pending_block[number] = match.group(2)
                continue

            flush_block()

            try:
                fields = next(csv.reader([line]))
            except csv.Error:
                continue

            if len(fields) < len(BASE_FIELDS):
                continue

            record = dict(zip(BASE_FIELDS, fields[:len(BASE_FIELDS)]))
            extra = {}

            for i, value in enumerate(fields[len(BASE_FIELDS):], start=13):
                col_name = column_map.get(i, f"Extra{i - 12}")
                extra[col_name] = value

            raw_index += 1
            record["extra"] = extra
            record["_raw_index"] = raw_index
            record["_event_time_int"] = safe_int(record.get("EventTime"))
            rows.append(record)

    return rows


# ---------------------------------------------------------------------------
# Participants and trials
# ---------------------------------------------------------------------------

def build_participant_numbers(rows):
    sessions = OrderedDict()

    for record in rows:
        key = session_key(record)
        sessions.setdefault(
            key,
            safe_int(record.get("ReceptionTime")) or 0,
        )

    ordered = sorted(
        sessions,
        key=lambda key: (sessions[key], repr(key)),
    )
    return {key: i + 1 for i, key in enumerate(ordered)}


def collect_custom_fields(rows):
    included = []
    excluded = set()

    for record in rows:
        for key, value in record.get("extra", {}).items():
            if is_empty(value):
                continue
            if is_pii_field(key):
                excluded.add(key)
            elif key not in included:
                included.append(key)

    return included, sorted(excluded)


def build_trials(rows):
    trials = OrderedDict()

    for record in rows:
        key = (
            session_key(record),
            record.get("OrderNumber", ""),
            record.get("Label", ""),
        )

        trial = trials.setdefault(key, {
            "SessionKey": session_key(record),
            "OrderNumber": record.get("OrderNumber", ""),
            "Label": decode_text(record.get("Label", "")),
            "Rows": [],
            "Sentence": None,
            "Custom": OrderedDict(),
        })

        trial["Rows"].append(record)

        for field, value in record.get("extra", {}).items():
            if is_empty(value) or is_pii_field(field):
                continue

            if is_sentence_like(field):
                if trial["Sentence"] is None:
                    trial["Sentence"] = decode_text(value)
            else:
                trial["Custom"].setdefault(field, decode_text(value))

    # Sort each trial by EventTime when possible, falling back to raw row order.
    for trial in trials.values():
        trial["Rows"].sort(
            key=lambda r: (
                r["_event_time_int"] is None,
                r["_event_time_int"] if r["_event_time_int"] is not None else r["_raw_index"],
                r["_raw_index"],
            )
        )

    return list(trials.values())


# ---------------------------------------------------------------------------
# SPR analysis and post-SPR extraction
# ---------------------------------------------------------------------------

def reading_time_from_record(record):
    for key, value in record.get("extra", {}).items():
        if is_reading_time_like(key):
            rt = safe_int(value)
            if rt is not None:
                return rt
    return None


def spr_segments(trial):
    segments = []

    for record in trial["Rows"]:
        if not event_is_spr(record):
            continue

        rt = reading_time_from_record(record)
        if rt is None:
            continue

        seg_index = safe_int(record.get("Parameter"))
        segments.append({
            "record": record,
            "index": seg_index,
            "text": decode_text(record.get("Value", "")),
            "rt": rt,
            "event_time": record["_event_time_int"],
            "raw_index": record["_raw_index"],
        })

    # Prefer segment number when available; retain raw/time order as tie breaker.
    segments.sort(
        key=lambda s: (
            s["index"] is None,
            s["index"] if s["index"] is not None else s["raw_index"],
            s["raw_index"],
        )
    )
    return segments


def final_spr_segment(trial):
    segments = spr_segments(trial)
    if not segments:
        return None, []

    # Final segment means the last segment in the SPR sequence, not merely the
    # last raw event carrying a DashedSentence-like label.
    return segments[-1], segments


def record_occurs_after_final_spr(record, final_segment):
    """
    Prefer absolute EventTime ordering. If EventTime is missing, fall back to
    raw file order. Exclude the final SPR record itself.
    """
    final_record = final_segment["record"]
    if record is final_record:
        return False

    event_time = record["_event_time_int"]
    final_time = final_segment["event_time"]

    if event_time is not None and final_time is not None:
        if event_time > final_time:
            return True
        if event_time < final_time:
            return False
        # Equal timestamps: preserve raw log order.
        return record["_raw_index"] > final_record["_raw_index"]

    return record["_raw_index"] > final_record["_raw_index"]


def get_post_spr_events(trial, final_segment):
    events = [
        record for record in trial["Rows"]
        if record_occurs_after_final_spr(record, final_segment)
    ]

    # We want every post-SPR event, including controller/scale/log events.
    return events


def reconstruct_sentence_from_segments(segments):
    return " ".join(
        str(seg["text"]).strip()
        for seg in segments
        if str(seg["text"]).strip()
    )


# ---------------------------------------------------------------------------
# Diagnostic: what does SPR EventTime represent?
# ---------------------------------------------------------------------------

def diagnostic_counts(trials):
    checked = 0
    start_matches = 0
    end_matches = 0
    examples = []

    for trial in trials:
        segments = spr_segments(trial)
        if len(segments) < 2:
            continue

        for earlier, later in zip(segments, segments[1:]):
            if earlier["event_time"] is None or later["event_time"] is None:
                continue

            gap = later["event_time"] - earlier["event_time"]
            diff_earlier = abs(gap - earlier["rt"])
            diff_later = abs(gap - later["rt"])

            checked += 1
            if diff_earlier <= diff_later:
                start_matches += 1
            else:
                end_matches += 1

            if len(examples) < 10:
                examples.append([
                    gap,
                    earlier["rt"],
                    later["rt"],
                    diff_earlier,
                    diff_later,
                ])

    return checked, start_matches, end_matches, examples


# ---------------------------------------------------------------------------
# Excel output
# ---------------------------------------------------------------------------

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(
    start_color="4472C4",
    end_color="4472C4",
    fill_type="solid",
)
BODY_FONT = Font(name="Arial")


def unique_headers(headers):
    seen = defaultdict(int)
    result = []

    for header in headers:
        name = str(header) if header else "Unnamed"
        seen[name] += 1
        result.append(
            name if seen[name] == 1 else f"{name}_{seen[name]}"
        )

    return result


def write_sheet(wb, title, headers, data_rows):
    ws = wb.create_sheet(title=title[:31])
    headers = unique_headers(headers)
    ws.append(headers)

    for row in data_rows:
        ws.append(row)

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top")

    for col_cells in ws.columns:
        values = [c.value for c in col_cells if c.value is not None]
        max_len = max((len(str(v)) for v in values), default=8)
        ws.column_dimensions[
            get_column_letter(col_cells[0].column)
        ].width = min(max_len + 2, 45)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build_workbook(rows):
    participant_number = build_participant_numbers(rows)
    trials = build_trials(rows)
    all_custom_fields, excluded_fields = collect_custom_fields(rows)

    # Retained fields appearing anywhere in the source. This keeps the schema
    # project-independent and makes Variables transparent.
    custom_fields = [
        field for field in all_custom_fields
        if not is_sentence_like(field)
    ]

    wb = Workbook()
    wb.remove(wb.active)

    trial_rows = []
    post_event_rows = []
    time_like_rows = []

    spr_trial_count = 0

    for trial in trials:
        final_seg, segments = final_spr_segment(trial)
        if final_seg is None:
            continue

        spr_trial_count += 1
        pnum = participant_number[trial["SessionKey"]]

        sentence = (
            trial["Sentence"]
            or reconstruct_sentence_from_segments(segments)
        )

        final_event_time = final_seg["event_time"]
        final_rt = final_seg["rt"]
        estimated_final_end = (
            final_event_time + final_rt
            if final_event_time is not None and final_rt is not None
            else None
        )

        post_events = get_post_spr_events(trial, final_seg)

        post_event_times = [
            record["_event_time_int"]
            for record in post_events
            if record["_event_time_int"] is not None
        ]
        first_post_time = min(post_event_times) if post_event_times else None
        last_post_time = max(post_event_times) if post_event_times else None

        trial_base = [
            pnum,
            safe_number_for_excel(trial["OrderNumber"]),
            trial["Label"],
            sentence,
        ]
        trial_base += [
            trial["Custom"].get(field, "")
            for field in custom_fields
        ]

        trial_rows.append(
            trial_base + [
                len(segments),
                final_seg["index"] if final_seg["index"] is not None else "",
                final_seg["text"],
                final_event_time if final_event_time is not None else "",
                final_rt if final_rt is not None else "",
                estimated_final_end if estimated_final_end is not None else "",
                len(post_events),
                first_post_time if first_post_time is not None else "",
                (
                    first_post_time - final_event_time
                    if first_post_time is not None and final_event_time is not None
                    else ""
                ),
                (
                    first_post_time - estimated_final_end
                    if first_post_time is not None and estimated_final_end is not None
                    else ""
                ),
                last_post_time if last_post_time is not None else "",
                (
                    last_post_time - final_event_time
                    if last_post_time is not None and final_event_time is not None
                    else ""
                ),
                (
                    last_post_time - estimated_final_end
                    if last_post_time is not None and estimated_final_end is not None
                    else ""
                ),
            ]
        )

        # Long format: EVERY post-SPR event.
        for event_index, record in enumerate(post_events, start=1):
            event_time = record["_event_time_int"]

            event_base = [
                pnum,
                safe_number_for_excel(trial["OrderNumber"]),
                trial["Label"],
                sentence,
            ]
            event_base += [
                trial["Custom"].get(field, "")
                for field in custom_fields
            ]

            event_custom = record.get("extra", {})

            post_event_rows.append(
                event_base + [
                    event_index,
                    record.get("Controller", ""),
                    record.get("PennElementType", ""),
                    record.get("PennElementName", ""),
                    record.get("Parameter", ""),
                    decode_text(record.get("Value", "")),
                    event_time if event_time is not None else "",
                    (
                        event_time - final_event_time
                        if event_time is not None and final_event_time is not None
                        else ""
                    ),
                    (
                        event_time - estimated_final_end
                        if event_time is not None and estimated_final_end is not None
                        else ""
                    ),
                ] + [
                    (
                        ""
                        if is_pii_field(field)
                        else decode_text(event_custom.get(field, ""))
                    )
                    for field in custom_fields
                ]
            )

            # Automatically discover time-like numeric custom fields on this event.
            for field, value in event_custom.items():
                if (
                    is_pii_field(field)
                    or is_empty(value)
                    or not is_time_like_field(field)
                ):
                    continue

                numeric_value = safe_number(value)
                if numeric_value is None:
                    continue

                time_like_rows.append([
                    pnum,
                    safe_number_for_excel(trial["OrderNumber"]),
                    trial["Label"],
                    sentence,
                    event_index,
                    record.get("Controller", ""),
                    record.get("PennElementType", ""),
                    record.get("PennElementName", ""),
                    record.get("Parameter", ""),
                    event_time if event_time is not None else "",
                    field,
                    numeric_value,
                ])

    trials_headers = [
        "Participant_Number",
        "OrderNumber",
        "Label",
        "Sentence",
    ] + custom_fields + [
        "N_SPR_Segments",
        "Final_SPR_Segment_Number",
        "Final_SPR_Segment_Text",
        "Final_SPR_EventTime",
        "Final_SPR_RT_ms",
        "Estimated_Final_SPR_EndTime",
        "N_Post_SPR_Events",
        "First_Post_SPR_EventTime",
        "Delta_FinalSPREvent_to_FirstPostEvent_ms",
        "Delta_EstimatedSPREnd_to_FirstPostEvent_ms",
        "Last_Post_SPR_EventTime",
        "Delta_FinalSPREvent_to_LastPostEvent_ms",
        "Delta_EstimatedSPREnd_to_LastPostEvent_ms",
    ]

    post_headers = [
        "Participant_Number",
        "OrderNumber",
        "Label",
        "Sentence",
    ] + custom_fields + [
        "Post_Event_Index",
        "Controller",
        "PennElementType",
        "PennElementName",
        "Parameter",
        "Value",
        "EventTime",
        "Delta_from_Final_SPR_EventTime_ms",
        "Delta_from_Estimated_Final_SPR_End_ms",
    ] + [
        f"EventVar_{field}" for field in custom_fields
    ]

    time_headers = [
        "Participant_Number",
        "OrderNumber",
        "Label",
        "Sentence",
        "Post_Event_Index",
        "Controller",
        "PennElementType",
        "PennElementName",
        "Parameter",
        "EventTime",
        "Time_Field_Name",
        "Time_Field_Value",
    ]

    trial_rows.sort(key=sort_trial_row)
    post_event_rows.sort(key=sort_post_event_row)
    time_like_rows.sort(key=sort_time_row)

    write_sheet(wb, "Trials", trials_headers, trial_rows)
    write_sheet(wb, "Post_SPR_Events", post_headers, post_event_rows)
    write_sheet(wb, "Time_Like_Fields", time_headers, time_like_rows)

    variable_rows = []
    for field in custom_fields:
        variable_rows.append([
            field,
            "Included",
            "Retained automatically as a non-personal custom variable",
            "Yes" if is_time_like_field(field) else "No",
        ])
    for field in excluded_fields:
        variable_rows.append([
            field,
            "Excluded",
            "Suppressed by privacy filter; values were not exported",
            "No",
        ])

    write_sheet(
        wb,
        "Variables",
        ["Variable", "Status", "Reason", "Looks_Time_Like"],
        variable_rows,
    )

    checked, start_matches, end_matches, examples = diagnostic_counts(trials)

    interpretation = (
        "Insufficient data"
        if checked == 0
        else (
            "EventTime most likely marks segment START; use deltas from Estimated_Final_SPR_End."
            if start_matches > end_matches
            else (
                "EventTime most likely marks segment END; use deltas from Final_SPR_EventTime."
                if end_matches > start_matches
                else "Diagnostic tied; inspect raw events manually."
            )
        )
    )

    diagnostic_rows = [
        ["Segment_pairs_checked", checked],
        ["Gap_closer_to_earlier_segment_RT", start_matches],
        ["Gap_closer_to_later_segment_RT", end_matches],
        ["Interpretation", interpretation],
        [],
        ["Example_Gap_ms", "Earlier_Segment_RT", "Later_Segment_RT",
         "AbsDiff_to_Earlier_RT", "AbsDiff_to_Later_RT"],
    ]
    diagnostic_rows.extend(examples)

    # Diagnostic is easier to read without the generic header writer because it
    # intentionally contains two small sections.
    ws = wb.create_sheet("Diagnostic")
    for row in diagnostic_rows:
        ws.append(row)
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 24
    ws.column_dimensions["D"].width = 24
    ws.column_dimensions["E"].width = 24

    return wb, {
        "participants": len(participant_number),
        "all_trials": len(trials),
        "spr_trials": spr_trial_count,
        "post_events": len(post_event_rows),
        "time_like_values": len(time_like_rows),
        "custom_fields": len(custom_fields),
        "excluded_fields": len(excluded_fields),
    }


def safe_number_for_excel(value):
    text = str(value).strip()
    if not text:
        return ""
    try:
        if re.fullmatch(r"[-+]?\d+", text):
            return int(text)
        return float(text)
    except ValueError:
        return value


def sort_trial_row(row):
    return (
        row[0] if isinstance(row[0], int) else 0,
        row[1] if isinstance(row[1], (int, float)) else float("inf"),
        str(row[2]),
    )


def sort_post_event_row(row):
    # Post_Event_Index sits after participant/order/label/sentence/custom fields,
    # so use participant/order first and preserve construction order otherwise.
    return (
        row[0] if isinstance(row[0], int) else 0,
        row[1] if isinstance(row[1], (int, float)) else float("inf"),
        str(row[2]),
    )


def sort_time_row(row):
    return (
        row[0] if isinstance(row[0], int) else 0,
        row[1] if isinstance(row[1], (int, float)) else float("inf"),
        row[4] if isinstance(row[4], int) else 0,
        str(row[10]),
    )


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "Extract every logged event after the final SPR segment of each "
            "PCIbex/PennController trial and calculate post-SPR timing."
        )
    )
    parser.add_argument(
        "input_file",
        type=Path,
        help="Raw PCIbex/PennController results CSV file",
    )
    parser.add_argument(
        "output_file",
        nargs="?",
        type=Path,
        default=Path(DEFAULT_OUTPUT_NAME),
        help=f"Output .xlsx file (default: {DEFAULT_OUTPUT_NAME})",
    )
    parser.add_argument(
        "--version",
        action="version",
        version=f"%(prog)s {VERSION}",
    )
    return parser.parse_args()


def main():
    args = parse_args()

    if not args.input_file.is_file():
        print(f"Error: input file not found: {args.input_file}", file=sys.stderr)
        sys.exit(1)

    if args.output_file.suffix.lower() != ".xlsx":
        print("Error: output file must end in .xlsx", file=sys.stderr)
        sys.exit(1)

    rows = parse_results_file(args.input_file)
    if not rows:
        print(
            "Error: no valid PCIbex/PennController data rows were found.",
            file=sys.stderr,
        )
        sys.exit(1)

    workbook, stats = build_workbook(rows)

    args.output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output_file)

    print(f"Participants (anonymised): {stats['participants']}")
    print(f"Trials reconstructed: {stats['all_trials']}")
    print(f"Trials containing SPR: {stats['spr_trials']}")
    print(f"Post-SPR events extracted: {stats['post_events']}")
    print(f"Time-like custom values found after SPR: {stats['time_like_values']}")
    print(f"Custom variables retained: {stats['custom_fields']}")
    print(f"Personal variables excluded: {stats['excluded_fields']}")
    print(f"Saved: {args.output_file}")


if __name__ == "__main__":
    main()
